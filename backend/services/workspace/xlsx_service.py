"""XLSX import/export helpers for MAC Analyzer Pro Web."""

from io import BytesIO
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill


def read_xlsx(content: bytes, sheet_name: str | None = None) -> dict[str, Any]:
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    sheet = workbook[sheet_name] if sheet_name else workbook.active
    rows = [list(row) for row in sheet.iter_rows(values_only=True)]
    if not rows:
        return {"sheet": sheet.title, "headers": [], "rows": []}
    headers = [str(value or f"Column {index + 1}") for index, value in enumerate(rows[0])]
    return {"sheet": sheet.title, "headers": headers, "rows": [["" if value is None else str(value) for value in row] for row in rows[1:]]}


def read_xls(content: bytes, sheet_name: str | None = None) -> dict[str, Any]:
    try:
        import xlrd
    except ImportError as error:
        raise ValueError("Legacy .xls import requires xlrd. Install requirements-web.txt or save the file as .xlsx.") from error

    workbook = xlrd.open_workbook(file_contents=content)
    sheet = workbook.sheet_by_name(sheet_name) if sheet_name else workbook.sheet_by_index(0)
    if sheet.nrows == 0:
        return {"sheet": sheet.name, "headers": [], "rows": []}
    raw_headers = sheet.row_values(0)
    headers = [str(value or f"Column {index + 1}") for index, value in enumerate(raw_headers)]
    rows = [
        ["" if value is None else str(value) for value in sheet.row_values(row_index)]
        for row_index in range(1, sheet.nrows)
    ]
    return {"sheet": sheet.name, "headers": headers, "rows": rows}


def export_xlsx(devices: list[dict[str, Any]], columns: list[tuple[str, str]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "MAC Analyzer"
    sheet.append([title for _, title in columns])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0F766E")
    for device in devices:
        sheet.append([device.get(key, "") for key, _ in columns])
    sheet.freeze_panes = "A2"
    for column in sheet.columns:
        letter = column[0].column_letter
        sheet.column_dimensions[letter].width = min(42, max(12, max(len(str(cell.value or "")) for cell in column) + 2))
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
